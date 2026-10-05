import os
import shutil
import subprocess

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest

from tests.test_export import REFERENCE, parse_reference
from transcript_studio.export_formats import export, layout, to_html
from transcript_studio.sounds import insert_tag


@pytest.fixture
def session():
    s = parse_reference(REFERENCE.read_text(encoding="utf-8"))
    for p in s.paragraphs:  # give sentences times so a tag can be placed
        for i, sent in enumerate(p.sentences):
            sent.start_time, sent.end_time = p.start_time + i, p.start_time + i + 0.9
    s.speaker_names = {"1": "Host"}
    s.info = {"audio": {"file": "talk.m4a", "duration": 852.3}, "transcription": {"engine": "whisper.cpp"}}
    insert_tag(s, 2.5, 2.7, "sneeze")
    return s


def _pdf_text(path):
    if not shutil.which("pdftotext"):
        pytest.skip("pdftotext not installed")
    return subprocess.run(["pdftotext", "-layout", str(path), "-"], capture_output=True, text=True).stdout


@pytest.mark.parametrize("details", [True, False])
def test_pdf_export(session, tmp_path, details):
    out = export(session, tmp_path / "t.pdf", details=details)
    text = _pdf_text(out)
    assert "NASA Podcast Ep 434 Excerpt Transcript" in text
    assert "Host" in text and "(sneeze)" in text and "Legend:" in text
    assert ("File details" in text) == details and ("Transcription details" in text) == details
    assert "19." in text  # every paragraph made it


@pytest.mark.parametrize("details", [True, False])
def test_docx_export(session, tmp_path, details):
    from docx import Document
    out = export(session, tmp_path / "t.docx", details=details)
    d = Document(str(out))
    heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert heads[0] == "NASA Podcast Ep 434 Excerpt Transcript"
    assert "2. ● Host — 00:00:00" in heads and "1. ● 00:00:00" in heads
    assert len(d.tables) == (2 if details else 0)
    if details:
        assert d.tables[0].cell(0, 0).text == "File" and d.tables[0].cell(0, 1).text == "talk.m4a"
        assert d.tables[0].cell(0, 1).paragraphs[0].runs[0].font.name == "Courier New"
    runs = [r for p in d.paragraphs for r in p.runs]
    sneeze = [r for r in runs if r.text == "(sneeze)"]
    assert sneeze and sneeze[0].italic
    low_text = session.paragraphs[2].sentences[0].text[:12]  # flagged low-confidence; the (sneeze) tag splits the rest
    low = [r for r in runs if low_text in r.text][0]
    assert str(low.font.color.rgb) == "DC2626"  # low-confidence sentences stay red


def test_markdown_via_dispatcher_matches_legacy(session, tmp_path):
    from transcript_studio.export import to_markdown
    out = export(session, tmp_path / "t.md", details=False)
    assert out.read_text(encoding="utf-8") == to_markdown(session, details=False)
    with pytest.raises(ValueError):
        export(session, tmp_path / "t.txt")


def test_html_uses_coloured_dots_not_emoji(session):
    h = to_html(session, details=False)
    assert "🔵" not in h and "●" in h and "<i>(sneeze)</i>" in h
    legend, blocks = layout(session)
    assert legend[0] == ("#1d4ed8", "Host") and len(blocks) == 19
